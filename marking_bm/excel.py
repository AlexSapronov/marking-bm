from io import BytesIO
import re
import zipfile
import openpyxl
import xlsxwriter

HEADERS = ["PN", "GTIN", "Кол-во в 1 этикетке", "Кол-во этикеток", "Код ТН ВЭД"]


def valid_gtin(s):
    return (
        bool(re.fullmatch("[0-9]{14}", s))
        and int(s[-1])
        == (
            10
            - sum(int(c) * (3 if i % 2 == 0 else 1) for i, c in enumerate(s[:13])) % 10
        )
        % 10
    )


def validate_rows(rows):
    errors, warnings, merged, pns, gtins = [], [], {}, {}, {}
    for n, row in enumerate(rows, 2):
        try:
            if not isinstance(row, (list, tuple)) or len(row) != 5:
                raise ValueError("Нужно пять столбцов")
            pn, gtin, quantity, labels, tnved = row
            if (
                not isinstance(pn, str)
                or not pn.strip()
                or len(pn.strip()) > 120
                or any(ord(c) < 32 for c in pn)
            ):
                raise ValueError("PN: нужен текст от 1 до 120 символов")
            pn = pn.strip()
            if not isinstance(gtin, str) or not valid_gtin(gtin):
                raise ValueError(
                    "GTIN: 14 цифр в текстовой ячейке, с верной контрольной цифрой"
                )
            if not isinstance(tnved, str) or not re.fullmatch("[0-9]{10}", tnved):
                raise ValueError("ТН ВЭД: 10 цифр в текстовой ячейке")
            for field, value, limit in [
                ("Штук в этикетке", quantity, 99999999),
                ("Этикеток", labels, 100000),
            ]:
                if type(value) is not int or not 1 <= value <= limit:
                    raise ValueError(f"{field}: целое число от 1 до {limit}")
            if pn in pns and pns[pn] != gtin:
                raise ValueError("Один PN соответствует разным GTIN")
            if gtin in gtins and gtins[gtin] != tnved:
                raise ValueError("Один GTIN соответствует разным ТН ВЭД")
            pns[pn], gtins[gtin] = gtin, tnved
            key = (pn, gtin, quantity, tnved)
            merged[key] = merged.get(key, 0) + labels
            if len(pn) > 35:
                warnings.append(
                    f"Строка {n}: длинный PN; проверьте размещение на этикетке"
                )
        except ValueError as e:
            errors.append({"row": n, "message": str(e)})
    if not merged and not errors:
        errors.append({"row": 2, "message": "Добавьте хотя бы одну строку"})
    if sum(merged.values()) > 100000:
        errors.append({"row": 0, "message": "Больше 100000 этикеток в одном документе"})
    return {
        "lines": [
            dict(pn=k[0], gtin=k[1], quantity=k[2], tnved=k[3], labels=v)
            for k, v in merged.items()
        ],
        "errors": errors,
        "warnings": list(dict.fromkeys(warnings)),
    }


def import_file(content):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("Размер Excel превышает 10 МБ")
    try:
        with zipfile.ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024:
                raise ValueError("Слишком большой распакованный Excel")
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
        ws = book.worksheets[0]
        if ws.max_row > 20001 or ws.max_column != 5:
            raise ValueError("Ожидается до 20000 строк и ровно пять столбцов")
        raw = list(ws.iter_rows())
        if [c.value for c in raw[0]] != HEADERS:
            raise ValueError("Используйте шаблон: неверный порядок столбцов")
        rows = []
        for row in raw[1:]:
            if all(c.value is None for c in row):
                rows.append(None)
            else:
                rows.append([c.value if c.data_type != "f" else None for c in row])
        result = validate_rows(
            [r if r is not None else ["", "", 0, 0, ""] for r in rows]
        )
        # Blank rows are ignored, while retaining physical row numbers in diagnostics.
        blank = {i + 2 for i, r in enumerate(rows) if r is None}
        result["errors"] = [e for e in result["errors"] if e["row"] not in blank]
        if not result["lines"] and not result["errors"]:
            result["errors"] = [{"row": 2, "message": "Нет заполненных строк"}]
        book.close()
        return result
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(
            "Не удалось прочитать XLSX. Используйте скачанный шаблон"
        ) from e


def xlsx(sheets):
    buf = BytesIO()
    wb = xlsxwriter.Workbook(
        buf, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}
    )
    head = wb.add_format({"bold": True, "bg_color": "#D8E9E7", "border": 1})
    text = wb.add_format({"num_format": "@"})
    for name, rows in sheets.items():
        ws = wb.add_worksheet(name)
        ws.freeze_panes(1, 0)
        ws.set_column(0, 4, 28, text)
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                if isinstance(v, str):
                    ws.write_string(r, c, v, head if r == 0 else text)
                elif v is not None:
                    ws.write_number(r, c, v, head if r == 0 else None)
    wb.close()
    return buf.getvalue()


def template():
    return xlsx({"Товары": [HEADERS]})
