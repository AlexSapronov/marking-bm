"""Legacy result_<invoice>.xlsx pack importer: one previously emitted KM per row."""
from collections import Counter
from io import BytesIO
from pathlib import Path
import re
import zipfile

import openpyxl

from .codes import GS, parse
from .excel import identifier, integer_value, pn_text, validate_rows


def invoice_from_filename(filename):
    name = Path(filename or "").name
    match = re.fullmatch(r"result_([A-Za-z0-9][A-Za-z0-9_-]{0,119})\.xlsx", name, re.I)
    if not match:
        raise ValueError("Файл должен называться result_НОМЕР.xlsx, например result_TLCWB.xlsx")
    return match[1]


def normalize_legacy_code(raw, quantity=None):
    if not isinstance(raw, str) or not raw:
        raise ValueError("Первый столбец должен содержать полный текст КМ")
    # Old pack exports can omit GS between GS1 fields and append AI 30.
    # The radio template has FIXED 20/4/44 lengths. Never search by "91",
    # "92" or "30" inside serial or crypto: those bytes may occur naturally.
    compact = re.fullmatch(
        r"01\d{14}21[\x21-\x7e]{20}91[\x21-\x7e]{4}"
        r"92[\x21-\x7e]{44}(?:30[0-9]{1,8})?",
        raw,
    )
    if compact:
        base = raw[:38] + GS + raw[38:44] + GS + raw[44:90]
        # AI 30 has no GS in old compact exports. Insert it only at the
        # known fixed boundary after the 44-byte cryptographic signature.
        raw = base + (GS + raw[90:] if len(raw) > 90 else "")
    code = parse(raw)
    if code.quantity is not None:
        if quantity is None:
            raise ValueError("КМ содержит AI 30: количество необходимо сверить с Excel")
        if code.quantity != quantity:
            raise ValueError(
                f"Количество в КМ (AI 30 = {code.quantity}) "
                f"не совпадает со столбцом D ({quantity})"
            )
        # codes.raw stores the cryptographic base WITHOUT AI 30. Later
        # with_quantity() adds the validated quantity exactly once.
        raw = code.full[:code.full.rfind(GS + "30")]
        code = parse(raw)
    return code


def parse_legacy_pack(content, filename):
    number = invoice_from_filename(filename)
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("Excel больше 10 МБ")
    if not zipfile.is_zipfile(BytesIO(content)):
        raise ValueError("Нужен настоящий файл XLSX")
    with zipfile.ZipFile(BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 100 or sum(entry.file_size for entry in entries) > 40 * 1024 * 1024:
            raise ValueError("Excel содержит слишком большие вложенные данные")
    try:
        book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            if "DataMatrix" not in book.sheetnames:
                raise ValueError("В архиве нужен лист DataMatrix")
            sheet = book["DataMatrix"]
            if sheet.max_row > 100000 or sheet.max_column > 20:
                raise ValueError("Слишком много строк или столбцов в Excel")
            items = []
            seen = set()
            for i, row in enumerate(sheet.iter_rows(min_row=1, max_col=5, values_only=True), 1):
                if not any(v is not None and str(v).strip() for v in row):
                    continue
                try:
                    raw, gtin_raw, pn_raw, quantity_raw, tnved_raw = row
                    quantity = integer_value(quantity_raw, "Штук в упаковке", 99999999)
                    code = normalize_legacy_code(raw, quantity)
                    gtin, _ = identifier(gtin_raw, 14, "GTIN")
                    if gtin != code.gtin:
                        raise ValueError("GTIN в столбце B отличается от GTIN внутри КМ")
                    pn = pn_text(pn_raw)
                    tnved, _ = identifier(tnved_raw, 10, "ТН ВЭД")
                    if code.cis in seen:
                        raise ValueError("Повторный КМ внутри файла")
                    seen.add(code.cis)
                    items.append({
                        "raw": code.full, "cis": code.cis, "serial": code.cis[18:38],
                        "gtin": gtin, "pn": pn, "quantity": quantity, "tnved": tnved,
                    })
                except ValueError as exc:
                    raise ValueError(f"Строка {i}: {exc}") from exc
        finally:
            book.close()
    except (OSError, zipfile.BadZipFile, KeyError) as exc:
        raise ValueError("Не удалось прочитать XLSX") from exc
    if not items:
        raise ValueError("В файле нет КМ")
    counts = Counter((x["pn"], x["gtin"], x["quantity"], x["tnved"]) for x in items)
    check = validate_rows([[pn, gtin, qty, count, tnved]
                           for (pn, gtin, qty, tnved), count in counts.items()])
    if check["errors"]:
        error = check["errors"][0]
        raise ValueError(f"Состав накладной: {error['message']}")
    return {"number": number, "items": items, "lines": check["lines"],
            "codes_count": len(items), "units": sum(x["quantity"] for x in items)}
