from io import BytesIO
import re
import sqlite3
from urllib.parse import unquote
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pytest

from conftest import GTIN, TN, add_doc, advance, client
from marking_bm.db import Store
from marking_bm.supplier_xls import TEMPLATE, codes_xls

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def data_rows(content):
    with ZipFile(BytesIO(content)) as z:
        strings = ET.fromstring(z.read('xl/sharedStrings.xml'))
        values = [''.join(t.itertext()) for t in strings]
        # Decode each escape once: a literal _x001D_ is stored _x005F_x001D_.
        values = [re.sub(r'_x([0-9a-fA-F]{4})_', lambda m: chr(int(m[1], 16)), v) for v in values]
        sheet = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        return [[values[int(c.find('s:v', NS).text)] for c in row]
                for row in sheet.find('s:sheetData', NS)]


def ready(env, number='INV-001', supplier='Supplier A'):
    app, service, _ = env
    doc = add_doc(service, number=number, supplier=supplier)
    service.order(doc, 'alice')
    advance(app)
    return doc


def enable(app, service, doc, enabled=True):
    supplier = service.get(doc, 'alice')['supplier_id']
    r = client(app, 'administrator', True).put(
        f'/api/admin/suppliers/{supplier}/codes-xls', json={'enabled': enabled})
    assert r.status_code == 200, r.text
    return supplier


def test_working_reference_contract_with_gs_crypto_and_text():
    raw = '01' + GTIN + '21' + '5' + 'A' * 19 + '\x1d91EE12\x1d92' + 'A' * 43 + '='
    pn = '=PN<&_x001D_'
    items = [{'raw': raw, 'gtin': GTIN, 'pn': pn}]
    content = codes_xls(items)
    assert content.startswith(b'PK\x03\x04')  # Working reference is OOXML with .xls extension.
    assert data_rows(content) == [[raw, GTIN, pn]]
    book = openpyxl.load_workbook(BytesIO(content))
    assert book.sheetnames == ['Sheet1']
    ws = book['Sheet1']
    assert (ws.max_row, ws.max_column) == (1, 3)
    assert all(ws.cell(1, col).data_type == 's' for col in (1, 2, 3))
    assert ws['B1'].value == GTIN and ws['C1'].data_type != 'f'
    assert ws['A1'].font.name == 'Calibri' and ws['A1'].font.sz == 11
    assert ws.column_dimensions['A'].width == 20.7109375
    assert ws.row_dimensions[1].height is None and ws.sheet_format.defaultRowHeight == 15
    with ZipFile(BytesIO(content)) as output, ZipFile(TEMPLATE) as template:
        assert output.namelist() == template.namelist()
        for name in template.namelist():
            if name not in ('xl/sharedStrings.xml', 'xl/worksheets/sheet1.xml'):
                assert output.read(name) == template.read(name)
        assert b'_x001D_91EE12_x001D_92' in output.read('xl/sharedStrings.xml')
        assert b'<f>' not in output.read('xl/worksheets/sheet1.xml')
    with pytest.raises(ValueError, match='без AI 30'):
        codes_xls([{**items[0], 'raw': raw + '\x1d305'}])
    with pytest.raises(ValueError, match='GTIN'):
        codes_xls([{**items[0], 'gtin': '00000000000000'}])


def test_template_has_no_original_codes_or_personal_properties():
    with ZipFile(TEMPLATE) as z:
        assert len(ET.fromstring(z.read('xl/sharedStrings.xml'))) == 0
        assert len(ET.fromstring(z.read('docProps/core.xml'))) == 0
        assert b'absPath' not in z.read('xl/workbook.xml')
        assert len(ET.fromstring(z.read('xl/worksheets/sheet1.xml')).find('s:sheetData', NS)) == 0


def test_setting_default_download_existing_ready_and_disable(env):
    app, service, _ = env
    doc = ready(env)
    url = f'/api/documents/{doc}/supplier-codes.xls'
    manager = client(app)
    before = service.get(doc, 'alice')
    assert not before['supplier_codes_xls_available']
    assert manager.get(url).status_code == 409
    enable(app, service, doc)
    assert service.get(doc, 'alice')['supplier_codes_xls_available']
    response = manager.get(url)
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == 'application/vnd.ms-excel'
    assert unquote(response.headers['content-disposition']).endswith("INV-001-Supplier A.xls")
    expected = service.store.all('SELECT raw,gtin,pn FROM codes WHERE doc_id=?', (doc,))
    assert sorted(data_rows(response.content)) == sorted([[r['raw'], r['gtin'], r['pn']] for r in expected])
    after = service.get(doc, 'alice')
    assert (before['status'], before['current_version'], before['artifacts'], before['codes']) == (
        after['status'], after['current_version'], after['artifacts'], after['codes'])
    enable(app, service, doc, False)
    assert manager.get(url).status_code == 409
    assert not service.get(doc, 'alice')['supplier_codes_xls_available']


def test_current_snapshot_after_repack_and_tampered_snapshot_blocked(env):
    app, service, _ = env
    doc = ready(env)
    enable(app, service, doc)
    raw = [['NEW-PN', GTIN, 9, 2, TN]]
    d = service.get(doc, 'alice')
    preview = service.preview(doc, 'alice', raw)
    service.save_lines(doc, 'alice', raw, d['revision'], preview['fingerprint'])
    name, content = service.supplier_codes_file(doc, 'alice')
    assert client(app).get(f'/api/documents/{doc}/supplier-codes.xls?version=1').status_code == 409
    assert len(data_rows(content)) == 2
    assert all(r[2] == 'NEW-PN' for r in data_rows(content))
    assert name == 'INV-001-Supplier A.xls'
    with service.store.transaction() as db:
        db.execute("UPDATE codes SET pn='OUTDATED' WHERE doc_id=?", (doc,))
    assert client(app).get(f'/api/documents/{doc}/supplier-codes.xls').status_code == 409


def test_permission_and_not_ready_and_stopped(env):
    app, service, _ = env
    doc = add_doc(service)
    supplier = enable(app, service, doc)
    url = f'/api/documents/{doc}/supplier-codes.xls'
    assert client(app).get(url).status_code == 409
    assert client(app).put(f'/api/admin/suppliers/{supplier}/codes-xls', json={'enabled': False}).status_code == 403
    service.order(doc, 'alice')
    advance(app)
    assert client(app, 'bob').get(url).status_code == 403
    admin = client(app, 'administrator', True)
    for role in ('logistics', 'supervisor'):
        admin.post('/api/admin/users', json={'username': role, 'role': role})
        assert client(app, role).get(url).status_code == 403
    assert admin.get(url).status_code == 200
    service.stop_document(doc, 'INV-1', 'alice')
    assert client(app).get(url).status_code == 409
    assert not service.get(doc, 'alice')['supplier_codes_xls_available']


def test_supplier_rename_and_merge_use_current_target_setting(env):
    app, service, _ = env
    doc = ready(env, number='INV/001')
    supplier = enable(app, service, doc)
    target = add_doc(service, supplier='Target', number='TARGET')
    target_id = service.get(target, 'alice')['supplier_id']
    admin = client(app, 'administrator', True)
    assert admin.post('/api/admin/suppliers', json={'id': supplier, 'name': 'Renamed'}).status_code == 200
    assert service.supplier_codes_file(doc, 'alice')[0] == 'INV_001-Renamed.xls'
    assert admin.post('/api/admin/suppliers', json={'id': supplier, 'name': 'Renamed', 'merge_into': target_id}).status_code == 200
    assert client(app).get(f'/api/documents/{doc}/supplier-codes.xls').status_code == 409
    enable(app, service, doc)
    assert service.supplier_codes_file(doc, 'alice')[0] == 'INV_001-Target.xls'


def test_old_supplier_migration_preserves_flags_on_restart(tmp_path):
    path = tmp_path / 'legacy.sqlite3'
    with sqlite3.connect(path) as db:
        db.execute('CREATE TABLE suppliers(id TEXT PRIMARY KEY,name TEXT NOT NULL UNIQUE,merged_into TEXT REFERENCES suppliers(id))')
        db.execute("INSERT INTO suppliers VALUES('old','Original',NULL)")
    store = Store(path)
    assert store.one("SELECT codes_xls FROM suppliers WHERE id='old'")['codes_xls'] == 0
    with store.transaction() as db:
        db.execute("UPDATE suppliers SET codes_xls=1 WHERE id='old'")
    assert Store(path).one("SELECT codes_xls FROM suppliers WHERE id='old'")['codes_xls'] == 1


def test_move_uses_new_target_and_never_source_history(env):
    app, service, _ = env
    source = ready(env, number='SOURCE')
    enable(app, service, source)
    before = data_rows(service.supplier_codes_file(source, 'alice')[1])
    target = service.create('alice', 'TARGET', 'Supplier A')
    line = service.get(source, 'alice')['lines'][0]['id']
    plan = service.move_lines_preview(source, target, [line], 'alice')
    service.move_lines(source, target, [line], plan['fingerprint'], 'alice')
    assert not service.get(source, 'alice')['supplier_codes_xls_available']
    assert client(app).get(f'/api/documents/{source}/supplier-codes.xls').status_code == 409
    name, content = service.supplier_codes_file(target, 'alice')
    assert name == 'TARGET-Supplier A.xls'
    assert data_rows(content) == before
